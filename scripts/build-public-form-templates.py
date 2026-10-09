#!/usr/bin/env python3
"""Generate the brand-free, print-ready files linked by the public tools."""
from pathlib import Path
import io
import re
import fitz
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "downloads"
OUT.mkdir(parents=True, exist_ok=True)
AFF_SOURCE = (ROOT / "components" / "AffidavitOfService.tsx").read_text()
FIELD_SOURCE = (ROOT / "components" / "FieldSheet.tsx").read_text()

# The published modes and their core wording must stay bound to the browser engine.
MODES = ["AFFIDAVIT OF SERVICE", "DECLARATION OF SERVICE", "AFFIDAVIT OF NON-SERVICE", "DECLARATION OF NON-SERVICE"]
for mode in MODES:
    assert mode in AFF_SOURCE, f"Generator is missing mode: {mode}"
for phrase in ("being duly sworn, depose and say", "declare under penalty of perjury under the laws of the State of", "Service Attempts (Physical):", "PERSON SERVED / ATTEMPTED:", "Subscribed and sworn to before me"):
    assert phrase in AFF_SOURCE, f"Generator wording drift: {phrase}"
assert "Field Attempts" in FIELD_SOURCE and "Service Execution Log" in FIELD_SOURCE

FIELD_LABELS = ["Court / Venue", "Case No.", "Plaintiff / Petitioner", "Defendant / Respondent", "Person Served / Attempted", "Server Name", "License #", "Documents", "Service Address", "Accepted By / Role", "Service Method", "Comments", "Notary State", "Notary County"]
FIELD_SHEET_LABELS = [("Agency / Server",724),("Phone",700),("Email",676),("Forwarding Client / Firm",652),("Client Phone / Contact",628),("Case No.",604),("Plaintiff",580),("Court / Venue",556),("Defendant",532),("Documents",508),("Due Date",484),("Special Instructions / Alerts",460),("PARTY TO SERVE",436),("Subject Phone",412),("Address 1 (Residence)",388),("Address 2 (Work / Business)",364),("Vehicle Info",340)]
FORM_NAMES = {
    "AFFIDAVIT OF SERVICE": "JLS-Blank-Affidavit-of-Service-v1.0",
    "DECLARATION OF SERVICE": "JLS-Blank-Declaration-of-Service-v1.0",
    "AFFIDAVIT OF NON-SERVICE": "JLS-Blank-Affidavit-of-Non-Service-v1.0",
    "DECLARATION OF NON-SERVICE": "JLS-Blank-Declaration-of-Non-Service-v1.0",
}


def write_pdf(mode: str, path: Path, fillable: bool = False):
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle(mode.title() + " — Blank Template")
    width = 612
    left, right = 42, 570
    c.setFont("Times-Roman", 10)
    c.setLineWidth(1.5)
    c.line(left, 742, right, 742)
    c.line(left, 678, right, 678)
    c.line(330, 742, 330, 678)
    def label(text, x, y, field=None, fw=150):
        c.setFont("Times-Bold", 9)
        c.drawString(x, y, text)
        lw = c.stringWidth(text, "Times-Bold", 9)
        if field and fillable:
            c.acroForm.textfield(name=field, tooltip=text, x=x+lw+7, y=y-4, width=fw, height=14, borderWidth=0.6, fontName="Helvetica", fontSize=9, forceBorder=True)
        else:
            c.setLineWidth(0.6)
            c.line(x+lw+7,y-4,min(right,x+lw+7+fw),y-4)
    label("PLAINTIFF / PETITIONER:", 48, 724, "plaintiff", 130)
    label("DEFENDANT / RESPONDENT:",48,700,"defendant",130)
    label("CASE NO.:",340,724,"case_no",135)
    label("PERSON SERVED / ATTEMPTED:",340,700,"recipient",86)
    c.setFont("Times-Bold", 12)
    c.drawCentredString(width/2, 654, mode)
    c.line(196, 651, 416, 651)
    declaration = mode.startswith("DECLARATION")
    nonservice = "NON-SERVICE" in mode
    c.setFont("Times-Roman", 10)
    if declaration:
        opening = "I, ___________________________, declare under penalty of perjury under the laws of the State of"
        continuation = "____________ and the United States of America that I am over 18, not a party to this action, and"
    else:
        opening = "I, ___________________________, being duly sworn, depose and say: I am over 18, not a party to this"
        continuation = "action, and I was authorized by law to serve the documents identified herein."
    c.drawString(left, 631, opening)
    c.drawString(left, 615, continuation)
    if declaration:
        c.drawString(left, 599, "was authorized by law to serve the documents identified herein.")
    label("Documents:",left,576,"documents",400)
    label("Service Address:",left,550,"service_address",383)
    c.setFont("Times-Bold",9)
    c.drawString(left,525,"SERVICE ATTEMPTS (PHYSICAL):")
    for idx,y in enumerate((502,480,458),1):
        c.setFont("Times-Roman",9)
        c.drawString(left+8,y,"Attempt %s — Date & Time:"%idx)
        if fillable:
            c.acroForm.textfield(name=f"attempt_{idx}",x=170,y=y-5,width=380,height=14,borderWidth=0.6,fontSize=9,forceBorder=True)
        else:c.line(170,y-5,right,y-5)
    c.setFont("Times-Bold",9)
    c.drawString(left,427,"COMMENTS / FIELD NOTES:")
    for y in (405,385,365):
        if fillable:
            c.acroForm.textfield(name=f"comments_{y}",x=left,y=y-4,width=right-left,height=14,borderWidth=0.6,fontSize=9,forceBorder=True)
        else: c.line(left,y-5,right,y-5)
    label("Method / Manner:",left,335,"method",395)
    if nonservice:
        execution = "After due search, careful inquiry, and diligent attempts at the address(es) above, I have been"
        following = "unable to effect process upon the recipient for the reasons recorded in the log above."
    else:
        execution = "I executed service on the recipient named above by the method and at the address recorded herein,"
        following = "delivering true and correct copies of the documents listed above as described in the comments."
    c.setFont("Times-Roman",10)
    c.drawString(left,302,execution)
    c.drawString(left,286,following)
    if declaration:
        c.drawString(left,268,"I state under penalty of perjury under the laws of the State of __________ that the foregoing")
        c.drawString(left,254,"is true and correct.")
    c.setLineWidth(0.7)
    c.line(left,190,292,190)
    c.setFont("Times-Roman",9)
    c.drawString(left,174,"Process Server / " + ("Declarant" if declaration else "Affiant"))
    label("Name:",left,155,"server",200)
    label("License # (if applicable):",left,136,"license",130)
    if declaration:
        label("Executed on:",320,226,"exec_date",118)
        label("At (city, state):",320,204,"exec_city",106)
    else:
        label("STATE OF",320,238,"notary_state",115)
        label("COUNTY OF",320,216,"notary_county",110)
        c.drawString(320,194,"Subscribed and sworn to before me on:")
        c.line(320,176,right,176)
        c.drawString(320,160,"Notary date")
        c.line(320,136,right,136)
        c.drawString(320,120,"Notary Public signature and seal")
    c.save()


def write_docx(mode: str, path: Path):
    d=Document(); sec=d.sections[0];sec.top_margin=Inches(.45);sec.bottom_margin=Inches(.45);sec.left_margin=sec.right_margin=Inches(.45)
    normal=d.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(10)
    t=d.add_table(rows=2, cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style="Table Grid"
    t.cell(0,0).text="PLAINTIFF / PETITIONER: ____________________"
    t.cell(1,0).text="DEFENDANT / RESPONDENT: __________________"
    t.cell(0,1).text="CASE NO.: ____________________"
    t.cell(1,1).text="PERSON SERVED / ATTEMPTED: _____________"
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(mode);r.bold=True;r.underline=True;r.font.size=Pt(12)
    if mode.startswith("DECLARATION"):
        d.add_paragraph("I, ___________________________, declare under penalty of perjury under the laws of the State of ____________ and the United States of America that I am over 18, not a party to this action, and was authorized by law to serve the documents identified herein.")
    else:
        d.add_paragraph("I, ___________________________, being duly sworn, depose and say: I am over 18, not a party to this action, and I was authorized by law to serve the documents identified herein.")
    for line in ["Documents: "+"_"*65,"Service Address: "+"_"*59]:d.add_paragraph(line)
    d.add_paragraph("SERVICE ATTEMPTS (PHYSICAL):").runs[0].bold=True
    t=d.add_table(rows=4,cols=2);t.style="Table Grid";t.cell(0,0).text="Attempt";t.cell(0,1).text="Date & Time"
    for i in range(1,4):t.cell(i,0).text=f"Attempt {i}";t.cell(i,1).text="__________________________"
    d.add_paragraph("COMMENTS / FIELD NOTES: "+"_"*55)
    d.add_paragraph("\n"+"_"*74)
    d.add_paragraph("Method / Manner: "+"_"*60)
    if "NON-SERVICE" in mode:
        d.add_paragraph("After due search, careful inquiry, and diligent attempts at the address(es) above, I have been unable to effect process upon the recipient for the reasons recorded in the log above.")
    else:
        d.add_paragraph("I executed service on the recipient named above by the method and at the address recorded herein, delivering true and correct copies of the documents listed above as described in the comments.")
    if mode.startswith("DECLARATION"):
        d.add_paragraph("I state under penalty of perjury under the laws of the State of __________ that the foregoing is true and correct.")
    t=d.add_table(rows=1,cols=2);t.style="Table Grid"
    t.cell(0,0).text="Signature: ________________________\nProcess Server / "+("Declarant" if mode.startswith("DECLARATION") else "Affiant")+"\nName: _____________________________\nLicense # (if applicable): ___________"
    if mode.startswith("DECLARATION"):
        t.cell(0,1).text="Executed on: ______________________\nAt (city, state): _____________________"
    else:
        t.cell(0,1).text="STATE OF ________________________\nCOUNTY OF ______________________\nSubscribed and sworn to before me on: ______\n_______________________________\nNotary Public signature and seal"
    d.save(path)

for mode in MODES:
    name=FORM_NAMES[mode]
    write_pdf(mode, OUT/(name+'.pdf'))
    write_docx(mode, OUT/(name+'.docx'))

# AcroForm remains fillable in the published file; bake a test *copy* for print verification.
write_pdf("AFFIDAVIT OF SERVICE",OUT/"JLS-Fillable-Affidavit-of-Service-v1.0.pdf",True)

# One-page field sheet aligned to the website's current field names and 4-attempt log.
fpath=OUT/"JLS-Blank-Field-Sheet-v1.0.pdf"

def field_pdf(path, fillable=False):
    c=canvas.Canvas(str(path),pagesize=letter);c.setTitle("PROCESS SERVER FIELD SHEET")
    c.setFont("Helvetica-Bold",15);c.drawCentredString(306,748,"PROCESS SERVER FIELD SHEET")
    labels=FIELD_SHEET_LABELS
    for label,y in labels:
        c.setFont("Helvetica-Bold",9);c.drawString(38,y,label+":")
        x=max(180,45+c.stringWidth(label+":","Helvetica-Bold",9))
        if fillable:
            c.acroForm.textfield(name=re.sub(r'[^a-z0-9]+','_',label.lower()).strip('_'),x=x,y=y-4,width=565-x,height=14,borderWidth=.6,fontSize=9,forceBorder=True)
        else:c.line(x,y-4,565,y-4)
    c.setFont("Helvetica-Bold",9);c.drawString(38,312,"FIELD ATTEMPTS & SERVICE EXECUTION LOG")
    for i,y in enumerate((288,267,246,225),1):
        c.setFont("Helvetica",9);c.drawString(42,y,f"{i}. Date / Time / Notes:")
        if fillable:c.acroForm.textfield(name=f"attempt_{i}",x=155,y=y-4,width=410,height=14,borderWidth=.6,fontSize=9,forceBorder=True)
        else:c.line(155,y-4,565,y-4)
    for label,y in [("Disposition (Served / Non-Service)",192),("Method / Accepted By / Role",168),("Assigned Server",144)]:
        c.setFont("Helvetica-Bold",9);c.drawString(38,y,label+":")
        x=max(245,45+c.stringWidth(label+":","Helvetica-Bold",9))
        if fillable:c.acroForm.textfield(name=re.sub(r'[^a-z0-9]+','_',label.lower()).strip('_'),x=x,y=y-4,width=565-x,height=14,borderWidth=.6,fontSize=9,forceBorder=True)
        else:c.line(x,y-4,565,y-4)
    c.setFont("Helvetica",9);c.drawString(38,112,"I certify that the field entries above reflect true and accurate service attempts.")
    c.line(38,82,340,82);c.drawString(38,68,"Process Server Signature & Date")
    c.save()

field_pdf(fpath)
field_pdf(OUT/"JLS-Fillable-Field-Sheet-v1.0.pdf",True)
d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.45);sec.left_margin=sec.right_margin=Inches(.5)
d.add_heading("PROCESS SERVER FIELD SHEET",0)
for label in [x for x,y in FIELD_SHEET_LABELS]:d.add_paragraph(f"{label}: "+"_"*52)
d.add_paragraph("FIELD ATTEMPTS & SERVICE EXECUTION LOG")
for i in range(1,5):d.add_paragraph(f"{i}. Date / Time / Notes: "+"_"*45)
for label in ("Disposition (Served / Non-Service)","Method / Accepted By / Role","Assigned Server"):
    d.add_paragraph(f"{label}: "+"_"*42)
d.add_paragraph("Process Server Signature & Date: "+"_"*35)
d.save(OUT/"JLS-Blank-Field-Sheet-v1.0.docx")

# Regression: verify every AcroForm field visibly survives a flatten-and-extract.
for path in [OUT/"JLS-Fillable-Affidavit-of-Service-v1.0.pdf",OUT/"JLS-Fillable-Field-Sheet-v1.0.pdf"]:
    reader=PdfReader(path);fields=reader.get_fields();assert fields and len(fields)>12
    values={name:f"TEST{i}" for i,name in enumerate(fields)}
    writer=PdfWriter();writer.append(reader);writer.update_page_form_field_values(writer.pages[0],values);writer.set_need_appearances_writer(False)
    buffer=io.BytesIO();writer.write(buffer);buffer.seek(0)
    pdf=fitz.open(stream=buffer.read(),filetype="pdf");pdf.bake(widgets=True)
    assert not any(list(p.widgets()) for p in pdf), f"Widgets remain: {path}"
    extracted="\n".join(p.get_text() for p in pdf)
    for value in values.values():assert value in extracted,(path,value)
    assert len(pdf)==1
    print(path.name,"fields",len(fields),"baked=0 widgets", "text verified")
print("Created",len(MODES)*2+4,"public files")
